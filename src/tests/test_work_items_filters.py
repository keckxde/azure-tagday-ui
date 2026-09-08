import unittest
import os
import sys
import tempfile

py_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if py_dir not in sys.path:
    sys.path.insert(0, py_dir)

from azure.azure_db import AzureDevOpsCache
from gui.backend import DevOpsBackend


class TestWorkItemsFilters(unittest.TestCase):
    def setUp(self):
        self.tmp_db = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
        self.tmp_db.close()
        self.cache = AzureDevOpsCache(self.tmp_db.name)

    def tearDown(self):
        if os.path.exists(self.tmp_db.name):
            try:
                os.remove(self.tmp_db.name)
            except Exception:
                pass

    def test_azure_db_work_item_iteration_fields(self):
        wi_id = 1001
        raw = {
            "id": wi_id,
            "fields": {
                "System.Title": "Implement feature X",
                "System.WorkItemType": "Feature",
                "System.State": "Active",
                "System.AssignedTo": {"displayName": "Alice Smith"},
                "System.ChangedDate": "2026-09-01T10:00:00Z",
                "System.IterationPath": "MyProject\\TeamA\\week-2630",
                "System.IterationId": 555,
                "System.AreaPath": "MyProject\\Area1",
            }
        }
        self.cache.save_work_item(
            wi_id,
            raw["fields"]["System.Title"],
            raw["fields"]["System.WorkItemType"],
            raw["fields"]["System.State"],
            "Alice Smith",
            raw["fields"]["System.ChangedDate"],
            raw,
            deleted=0
        )

        item = self.cache.get_work_item(wi_id)
        self.assertIsNotNone(item)
        self.assertEqual(item["iteration_path"], "MyProject\\TeamA\\week-2630")
        self.assertEqual(item["iteration_id"], 555)
        self.assertEqual(item["area_path"], "MyProject\\Area1")

        all_items = self.cache.get_all_work_items()
        self.assertEqual(len(all_items), 1)
        self.assertEqual(all_items[0]["iteration_path"], "MyProject\\TeamA\\week-2630")
        self.assertEqual(all_items[0]["iteration_id"], 555)
        self.assertEqual(all_items[0]["area_path"], "MyProject\\Area1")

    def test_backend_iteration_classification(self):
        # Create work items with different iteration hierarchies:
        # 1. Project root -> Unplanned
        # 2. Team area parent container -> Unplanned
        # 3. Dev Backlog -> Unplanned (contains backlog)
        # 4. Sprint week-2630 -> Planned
        # 5. Sprint week-2633 -> Planned
        items = [
            (2001, "Root item", "Task", "Active", "Dev 1", "2026-09-05T12:00:00Z", "ProjX"),
            (2002, "Team backlog item", "Task", "Active", "Dev 2", "2026-09-04T12:00:00Z", "ProjX\\TeamA"),
            (2003, "Dev backlog item", "Task", "Active", "Dev 1", "2026-09-03T12:00:00Z", "ProjX\\TeamA\\Dev Backlog"),
            (2004, "Planned sprint 1", "Bug", "Resolved", "Dev 3", "2026-09-02T12:00:00Z", "ProjX\\TeamA\\week-2630"),
            (2005, "Planned sprint 2", "Feature", "In Progress", "Dev 2", "2026-08-20T12:00:00Z", "ProjX\\TeamA\\week-2633"),
        ]

        for wi_id, title, type_str, state, user, date_str, iter_path in items:
            raw = {
                "id": wi_id,
                "fields": {
                    "System.Title": title,
                    "System.WorkItemType": type_str,
                    "System.State": state,
                    "System.AssignedTo": {"displayName": user},
                    "System.ChangedDate": date_str,
                    "System.IterationPath": iter_path,
                }
            }
            self.cache.save_work_item(wi_id, title, type_str, state, user, date_str, raw, deleted=0)

        # Setup backend with this cache database
        backend = DevOpsBackend()
        backend._cache_db = self.cache
        backend._db_path = self.tmp_db.name
        backend.refresh_all_data()

        work_items = backend.workItems
        self.assertEqual(len(work_items), 5)

        by_id = {wi["id"]: wi for wi in work_items}

        # Item 2001: Root path -> Unplanned
        self.assertFalse(by_id[2001]["is_iteration_planned"])

        # Item 2002: Team parent container -> Unplanned
        self.assertFalse(by_id[2002]["is_iteration_planned"])

        # Item 2003: Backlog leaf -> Unplanned
        self.assertFalse(by_id[2003]["is_iteration_planned"])

        # Item 2004: week-2630 -> Planned
        self.assertTrue(by_id[2004]["is_iteration_planned"])
        self.assertEqual(by_id[2004]["iteration_name"], "week-2630")

        # Item 2005: week-2633 -> Planned
        self.assertTrue(by_id[2005]["is_iteration_planned"])
        self.assertEqual(by_id[2005]["iteration_name"], "week-2633")

        # Test backend properties
        iterations = backend.workItemIterations
        self.assertEqual(iterations, ["week-2630", "week-2633"])

        assignees = backend.workItemAssignees
        self.assertEqual(assignees, ["Dev 1", "Dev 2", "Dev 3"])

    def test_export_work_items_to_excel(self):
        backend = DevOpsBackend()
        backend._work_items = [
            {
                "id": 3001,
                "title": "Implement Login",
                "type": "Requirement",
                "state": "Active",
                "assigned_to": "Alice",
                "iteration_path": "Project\\week-2633",
                "sprint_week_name": "week-2633",
                "deadline_str": "2026-08-14",
                "urgency_status": "overdue",
                "milestone_name": "DDQS Gate 1",
                "milestone_category": "Internal Process (DDQS)",
                "level1_display": "[10] Powertrain",
                "level1_pbs": "10",
                "level2_display": "[10.1] Battery Control",
                "level2_pbs": "10.1",
                "is_prio1": True,
                "is_grouped": True,
                "remaining_work": 6.0,
                "completed_work": 2.0,
                "changed_date": "2026-08-10T12:00:00Z",
                "tfs_url": "https://tfs.example.com/workitem/3001",
            },
            {
                "id": 3002,
                "title": "Fix Memory Leak",
                "type": "Bug",
                "state": "Closed",
                "assigned_to": "Bob",
                "iteration_path": "Project\\week-2634",
                "sprint_week_name": "week-2634",
                "deadline_str": "2026-08-21",
                "urgency_status": "completed",
                "milestone_name": "",
                "milestone_category": "",
                "level1_display": "",
                "level1_pbs": "",
                "level2_display": "",
                "level2_pbs": "",
                "is_prio1": False,
                "is_grouped": False,
                "remaining_work": 0.0,
                "completed_work": 4.0,
                "changed_date": "2026-08-15T15:30:00Z",
                "tfs_url": "https://tfs.example.com/workitem/3002",
            }
        ]

        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl is not installed in the test environment")

        with tempfile.TemporaryDirectory() as tmpdir:
            out_file = os.path.join(tmpdir, "exported_workitems.xlsx")
            res = backend.exportWorkItemsToExcel(file_path=out_file)
            self.assertTrue(res["success"])
            self.assertEqual(res["item_count"], 2)
            self.assertTrue(os.path.exists(out_file))

            wb = openpyxl.load_workbook(out_file)
            self.assertIn("Work Items", wb.sheetnames)
            ws = wb["Work Items"]

            # Header assertions
            self.assertEqual(ws.cell(row=1, column=1).value, "ID")
            self.assertEqual(ws.cell(row=1, column=2).value, "Type")
            self.assertEqual(ws.cell(row=1, column=3).value, "Title")
            self.assertEqual(ws.cell(row=1, column=4).value, "State")

            # Row 1 assertions
            self.assertEqual(ws.cell(row=2, column=1).value, 3001)
            self.assertEqual(ws.cell(row=2, column=2).value, "Requirement")
            self.assertEqual(ws.cell(row=2, column=3).value, "Implement Login")
            self.assertEqual(ws.cell(row=2, column=5).value, "Alice")
            self.assertEqual(ws.cell(row=2, column=10).value, "DDQS Gate 1")

            # Row 2 assertions
            self.assertEqual(ws.cell(row=3, column=1).value, 3002)
            self.assertEqual(ws.cell(row=3, column=2).value, "Bug")
            self.assertEqual(ws.cell(row=3, column=4).value, "Closed")
            self.assertEqual(ws.cell(row=3, column=5).value, "Bob")

    def test_work_items_tags_extraction_and_properties(self):
        backend = DevOpsBackend()
        backend._work_items = [
            {
                "id": 4001,
                "title": "Setup CI/CD Pipeline",
                "type": "Task",
                "state": "Active",
                "tags": "Target:DDQS-01; DevOps; Prio1",
                "tag_list": ["Target:DDQS-01", "DevOps", "Prio1"],
                "target_tags": ["DDQS-01"],
            },
            {
                "id": 4002,
                "title": "Database Migration",
                "type": "Requirement",
                "state": "Active",
                "tags": "Target:DDQS-01; Backend; Infra",
                "tag_list": ["Target:DDQS-01", "Backend", "Infra"],
                "target_tags": ["DDQS-01"],
            },
            {
                "id": 4003,
                "title": "Frontend Redesign",
                "type": "User Story",
                "state": "New",
                "tags": "Target:DDQS-02; Frontend; UI",
                "tag_list": ["Target:DDQS-02", "Frontend", "UI"],
                "target_tags": ["DDQS-02"],
            },
            {
                "id": 4004,
                "title": "Untagged Cleanup",
                "type": "Task",
                "state": "Closed",
                "tags": "",
                "tag_list": [],
                "target_tags": [],
            }
        ]

        # Test workItemTags property
        tags = backend.workItemTags
        self.assertIn("Backend", tags)
        self.assertIn("DevOps", tags)
        self.assertIn("Frontend", tags)
        self.assertIn("Infra", tags)
        self.assertIn("Prio1", tags)
        self.assertIn("Target:DDQS-01", tags)
        self.assertIn("Target:DDQS-02", tags)
        self.assertIn("UI", tags)

        # Test workItemTargetTags property
        target_tags = backend.workItemTargetTags
        self.assertIn("DDQS-01", target_tags)
        self.assertIn("DDQS-02", target_tags)

        # Test workItemTagCounts property
        counts = backend.workItemTagCounts
        self.assertEqual(counts["Target:DDQS-01"], 2)
        self.assertEqual(counts["Backend"], 1)
        self.assertEqual(counts["UI"], 1)

        # Test get_work_item_tags_summary()
        summary = backend.get_work_item_tags_summary()
        self.assertEqual(summary[0]["tag"], "Target:DDQS-01")
        self.assertEqual(summary[0]["count"], 2)
        self.assertTrue(summary[0]["is_target"])

        # Test target tag prefill in milestones
        milestones = backend.workItemMilestones
        self.assertIn("DDQS-01", milestones)
        self.assertIn("DDQS-02", milestones)


if __name__ == "__main__":
    unittest.main()


